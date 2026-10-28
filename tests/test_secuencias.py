#Secuencias didácticas: datos, vínculos, validación, fechas y exportación Word.
#Corre en un subproceso con una base temporal para no tocar datos reales.
import subprocess, sys, textwrap
from pathlib import Path

RAIZ= Path(__file__).resolve().parents[1]

SCRIPT= textwrap.dedent("""
    import sys, os
    from src.nucleo.bd import inicializar_bd, obtener_conexion
    inicializar_bd()
    from src.nucleo import secuencias as S

    conn= obtener_conexion(); c = conn.cursor()
    c.execute("INSERT INTO grados (nombre, orden) VALUES ('1ro', 1)"); g = c.lastrowid
    c.execute("INSERT INTO materias (codigo, nombre) VALUES ('LEN','Lengua Española')"); m = c.lastrowid
    c.execute("INSERT INTO materias (codigo, nombre) VALUES ('MAT','Matemática')"); m2 = c.lastrowid
    c.execute("INSERT INTO profesores (nombre, apellidos) VALUES ('Ana','Pérez')"); p = c.lastrowid
    c.execute("INSERT INTO competencias_especificas (materia_id, grado_id, codigo, descripcion) "
              "VALUES (?,?,'CE1','Comprende textos')", (m, g)); ce = c.lastrowid
    c.execute("INSERT INTO indicadores_logro (competencia_id, codigo, descripcion) "
              "VALUES (?,'IL1','Identifica la idea principal')", (ce,)); il = c.lastrowid
    c.execute("INSERT INTO tareas (materia_id, titulo, fecha_entrega) VALUES (?, 'Ensayo', '2026-11-20')", (m,)); t = c.lastrowid
    conn.commit(); conn.close()

    #--- catálogos
    assert [x["id"] for x in S.indicadores_de_competencias([ce])] == [il]
    assert S.indicadores_de_competencias([]) == []
    assert [x["id"] for x in S.tareas_de_materia(m)] == [t]

    #--- sesiones base y fechas
    base= S.sesiones_base(4)
    assert [s["momento"] for s in base] == ["Inicio", "Desarrollo", "Desarrollo", "Cierre"]
    assert [s["momento"] for s in S.sesiones_base(1)] == ["Desarrollo"]
    #2026-10-23 es viernes: salta fin de semana y el feriado del lunes
    assert S.distribuir_fechas("2026-10-23", 3, feriados=["2026-10-26"]) == ["2026-10-23", "2026-10-27", "2026-10-28"]
    assert S.distribuir_fechas("2026-10-26", 2, dias_semana=(0, 2)) == ["2026-10-26", "2026-10-28"]

    #--- validaciones de datos
    for malo in ({"titulo": "  "}, {"titulo": "X", "fecha_inicio": "2026-11-10", "fecha_fin": "2026-11-01"}):
        try:
            S.guardar_secuencia(malo); raise SystemExit("debió fallar")
        except ValueError:
            pass

    #--- crear completa
    ses= S.sesiones_base(3)
    ses[0].update(fecha="2026-11-02", actividad="Lluvia de ideas", evidencia="Mapa", instrumento="Lista de cotejo")
    ses[1].update(actividad="Lectura guiada", evidencia="Resumen")  #sin instrumento
    sid= S.guardar_secuencia(
        {"titulo": "Los textos narrativos", "materia_id": m, "grado_id": g, "profesor_id": p,
         "situacion_aprendizaje": "Los estudiantes leen cuentos de su comunidad.",
         "competencias_fundamentales": S.unir_competencias_fundamentales(["Comunicativa", "Ética y Ciudadana"]),
         "contenido_conceptos": "El cuento", "estado": "Estado inválido",
         "fecha_inicio": "2026-11-02", "fecha_fin": "2026-11-13"},
        competencias_ids= [ce, ce], indicadores_ids=[il], tareas_ids=[t], sesiones=ses)
    sec= S.obtener_secuencia(sid)
    assert sec["estado"] == "Borrador"                       #estado inválido -> Borrador
    assert sec["materia"] == "Lengua Española" and sec["grado"] == "1ro" and sec["profesor"] == "Ana Pérez"
    assert [x["id"] for x in sec["competencias"]] == [ce]    #duplicado ignorado
    assert [x["codigo"] for x in sec["indicadores"]] == ["IL1"]
    assert [x["titulo"] for x in sec["tareas"]] == ["Ensayo"]
    assert [s["orden"] for s in sec["sesiones"]] == [1, 2, 3]
    assert sec["competencias_fundamentales_lista"] == ["Comunicativa", "Ética y Ciudadana"]

    #--- validación pedagógica
    avisos= S.validar_secuencia(sec)
    assert any("Sesiones sin actividad descrita: 3" in a for a in avisos), avisos
    assert any("sin instrumento" in a and "2" in a for a in avisos), avisos
    assert not any("situación de aprendizaje" in a for a in avisos)
    assert S.validar_secuencia({"sesiones": []})  #vacía -> muchos avisos

    #--- avance
    assert S.progreso(sec["sesiones"]) == (0, 3, 0)
    S.marcar_sesion_realizada(sec["sesiones"][0]["id"])
    assert S.progreso(S.obtener_secuencia(sid)["sesiones"]) == (1, 3, 33)

    #--- editar sin tocar sesiones (sesiones=None las conserva) y reemplazando
    S.guardar_secuencia({"titulo": "Los textos narrativos II", "materia_id": m}, secuencia_id=sid,
                        competencias_ids= [ce])
    sec= S.obtener_secuencia(sid)
    assert sec["titulo"].endswith("II") and len(sec["sesiones"]) == 3 and sec["indicadores"] == []
    try:
        S.guardar_secuencia({"titulo": "x"}, secuencia_id=99999); raise SystemExit("debió fallar")
    except ValueError:
        pass

    #--- listado y filtros (incluye restricción tipo Profesor)
    sid2= S.guardar_secuencia({"titulo": "Fracciones", "materia_id": m2, "estado": "En curso"})
    assert {r["titulo"] for r in S.listar_secuencias()} == {"Los textos narrativos II", "Fracciones"}
    assert [r["titulo"] for r in S.listar_secuencias(materias_permitidas={m})] == ["Los textos narrativos II"]
    assert [r["titulo"] for r in S.listar_secuencias(materias_permitidas=set())] == []
    #un Profesor ve además lo que él creó aunque la materia no sea suya
    S.guardar_secuencia({"titulo": "Propia", "materia_id": m2, "profesor_id": p})
    assert "Propia" in {r["titulo"] for r in S.listar_secuencias(profesor_id=p, materias_permitidas={m})}
    assert [r["titulo"] for r in S.listar_secuencias(estado="En curso")] == ["Fracciones"]
    assert [r["titulo"] for r in S.listar_secuencias(texto="narrativos")] == ["Los textos narrativos II"]
    assert S.listar_secuencias(materia_id=m)[0]["total_sesiones"] == 3
    assert S.listar_secuencias(materia_id=m)[0]["sesiones_realizadas"] == 1

    #--- duplicar
    dup= S.obtener_secuencia(S.duplicar_secuencia(sid))
    assert dup["titulo"].endswith("(copia)") and dup["estado"] == "Borrador"
    assert dup["fecha_inicio"] is None and all(s["fecha"] is None and not s["realizada"] for s in dup["sesiones"])
    assert len(dup["sesiones"]) == 3 and [x["id"] for x in dup["competencias"]] == [ce]

    #--- exportar Word
    sec_full= S.guardar_secuencia({"titulo": "Completa", "materia_id": m, "grado_id": g,
        "situacion_aprendizaje": "Situación X", "contenido_conceptos": "C1", "estrategia": "ABP"},
        competencias_ids= [ce], indicadores_ids=[il], tareas_ids=[t], sesiones=ses)
    ruta= os.path.join(os.environ["HOME"], "sec.docx")
    S.exportar_docx(sec_full, ruta)
    from docx import Document
    d= Document(ruta)
    texto= "\\n".join(p.text for p in d.paragraphs)
    celdas= {cel.text for tb in d.tables for fila in tb.rows for cel in fila.cells}
    assert "Situación X" in texto and "CE1. Comprende textos" in texto and "ABP" in texto
    assert "Ensayo" in texto and "IL1: Identifica la idea principal" in texto
    assert {"Lluvia de ideas", "Lista de cotejo", "02/11/2026", "Lengua Española"} <= celdas, celdas
    try:
        S.exportar_docx(99999, ruta); raise SystemExit("debió fallar")
    except ValueError:
        pass

    #--- borrar secuencia limpia sesiones y vínculos (cascade), no toca tareas ni CE
    S.eliminar_secuencia(sid)
    assert S.obtener_secuencia(sid) is None
    conn= obtener_conexion()
    assert conn.execute("SELECT COUNT(*) FROM secuencia_sesiones WHERE secuencia_id=?", (sid,)).fetchone()[0] == 0
    assert conn.execute("SELECT COUNT(*) FROM tareas").fetchone()[0] == 1
    assert conn.execute("SELECT COUNT(*) FROM competencias_especificas").fetchone()[0] == 1
    #borrar la materia NO borra la secuencia (SET NULL): la planificación se conserva
    conn.execute("DELETE FROM materias WHERE id=?", (m2,)); conn.commit()
    assert S.obtener_secuencia(sid2)["materia_id"] is None
    conn.close()
    print("OK")
""")


def _correr(script, tmp_path):

    env= {"HOME": str(tmp_path), "XDG_DATA_HOME": str(tmp_path), "APPDATA": str(tmp_path), "USERPROFILE": str(tmp_path), "QT_QPA_PLATFORM": "offscreen",
           "PYTHONDONTWRITEBYTECODE": "1", "PYTHONIOENCODING": "utf-8", "PYTHONUTF8": "1", "PATH": "/usr/bin:/bin:/usr/local/bin"}
    return subprocess.run([sys.executable, "-c", script], cwd=RAIZ, env=env,
                          capture_output= True, text=True, timeout=180)


def test_secuencias_logica_y_exportacion(tmp_path):

    r= _correr(SCRIPT, tmp_path)
    assert r.returncode == 0 and "OK" in r.stdout, (r.stdout + r.stderr)[-2500:]


def test_competencias_fundamentales_coinciden_con_configuracion():

    from src.interfaz.configuracion import COMPETENCIAS_FUNDAMENTALES_RD
    from src.nucleo.secuencias import COMPETENCIAS_FUNDAMENTALES
    assert COMPETENCIAS_FUNDAMENTALES == COMPETENCIAS_FUNDAMENTALES_RD


SCRIPT_UI= textwrap.dedent("""
    import os, sys
    from PySide6.QtWidgets import QApplication, QWidget, QCheckBox
    from PySide6.QtCore import Qt, QDate
    app= QApplication(sys.argv)
    from src.nucleo.bd import inicializar_bd, obtener_conexion
    inicializar_bd()
    conn= obtener_conexion(); c = conn.cursor()
    c.execute("INSERT INTO grados (nombre, orden) VALUES ('2do', 2)"); g = c.lastrowid
    c.execute("INSERT INTO secciones (grado_id, nombre) VALUES (?, 'A')", (g,))
    c.execute("INSERT INTO materias (codigo, nombre) VALUES ('CS','Ciencias Sociales')"); m = c.lastrowid
    c.execute("INSERT INTO competencias_especificas (materia_id, grado_id, codigo, descripcion) "
              "VALUES (?,?,'CE1','Analiza hechos')", (m, g)); ce = c.lastrowid
    c.execute("INSERT INTO indicadores_logro (competencia_id, codigo, descripcion) VALUES (?,'IL1','Ubica en el tiempo')", (ce,))
    c.execute("INSERT INTO tareas (materia_id, titulo) VALUES (?, 'Línea de tiempo')", (m,))
    conn.commit(); conn.close()

    import src.modulos.secuencias as ui
    ui.mostrar_mensaje= lambda *a, **k: None          #sin diálogos modales en la prueba
    from src.nucleo import secuencias as S
    mw= QWidget(); mw.rol_actual = None; mw.id_usuario_actual = None

    modulo= ui.ModuloSecuenciasDidacticas(mw)
    assert modulo.tabla_secuencias.rowCount() == 0 and modulo.lbl_contador.text() == "0 secuencias"

    #editor nuevo: rellenar por la interfaz y guardar
    ed= ui.EditorSecuencia(mw)
    assert ed.tabla_sesiones.rowCount() == 3
    ed.in_titulo.setText("La independencia")
    ed.combo_materia.setCurrentIndex(ed.combo_materia.findData(m))
    ed.combo_grado.setCurrentIndex(ed.combo_grado.findData(g))
    assert ed.lista_competencias.count() == 1 and ed.lista_tareas.count() == 1 and ed.lista_indicadores.count() == 0
    ed.lista_competencias.item(0).setCheckState(Qt.CheckState.Checked)       #dispara la carga de indicadores
    assert ed.lista_indicadores.count() == 1
    ed.lista_indicadores.item(0).setCheckState(Qt.CheckState.Checked)
    ed.lista_tareas.item(0).setCheckState(Qt.CheckState.Checked)
    ed.lista_fundamentales.item(1).setCheckState(Qt.CheckState.Checked)
    ed.in_situacion.setPlainText("Investigan el 27 de febrero.")
    ed.tabla_sesiones.item(0, 2).setText("Lluvia de ideas")
    ed.in_fecha_inicio.setDate(QDate(2026, 11, 2))
    ed.tabla_sesiones.agregar_despues_de_la_actual()
    assert ed.tabla_sesiones.rowCount() == 4
    ed.tabla_sesiones.mover_actual(-1)
    ed.tabla_sesiones.selectRow(1); ed.tabla_sesiones.mover_actual(1)
    ed.tabla_sesiones.cellWidget(0, 8).findChild(QCheckBox).setChecked(True)
    ed._guardar()
    assert ed.result() == 1 and ed.secuencia_id
    sec= S.obtener_secuencia(ed.secuencia_id)
    assert sec["materia"] == "Ciencias Sociales" and sec["grado"] == "2do"
    assert [x["codigo"] for x in sec["competencias"]] == ["CE1"] and len(sec["indicadores"]) == 1
    assert [t["titulo"] for t in sec["tareas"]] == ["Línea de tiempo"]
    assert sec["competencias_fundamentales_lista"] == ["Comunicativa"]
    assert len(sec["sesiones"]) == 4 and S.progreso(sec["sesiones"])[0] == 1
    assert sec["sesiones"][0]["actividad"] == "Lluvia de ideas"

    #reabrir: debe volver a marcar todo lo guardado
    ed2= ui.EditorSecuencia(mw, ed.secuencia_id)
    assert ed2.in_titulo.text() == "La independencia"
    assert ed2.combo_materia.currentData() == m and ed2.combo_grado.currentData() == g
    assert ed2._marcados(ed2.lista_competencias) == [ce] and len(ed2._marcados(ed2.lista_indicadores)) == 1
    assert len(ed2._marcados(ed2.lista_tareas)) == 1
    assert ed2.tabla_sesiones.rowCount() == 4 and ed2.in_fecha_inicio.date() == QDate(2026, 11, 2)
    assert ed2.combo_seccion.currentData() is None
    assert ed2.lista_fundamentales.item(1).checkState() == Qt.CheckState.Checked
    #autocompletar fechas (sin diálogo): invoca la lógica con los mismos datos
    ses= ed2.tabla_sesiones.sesiones()
    assert ses[0]["realizada"] is True and ses[0]["actividad"] == "Lluvia de ideas"

    #lista principal y exportación
    modulo.refrescar()
    assert modulo.tabla_secuencias.rowCount() == 1 and modulo.tabla_secuencias.item(0, 0).text() == "La independencia"
    assert modulo.tabla_secuencias.item(0, 5).text() == "1/4 sesiones"
    modulo.tabla_secuencias.selectRow(0)
    modulo._duplicar(); assert modulo.tabla_secuencias.rowCount() == 2
    modulo.in_buscar.setText("zzz"); assert modulo.tabla_secuencias.rowCount() == 0
    modulo.in_buscar.setText(""); modulo.actualizar_tema()
    print("OK")
""")


def test_secuencias_interfaz(tmp_path):

    r= _correr(SCRIPT_UI, tmp_path)
    assert r.returncode == 0 and "OK" in r.stdout, (r.stdout + r.stderr)[-3000:]


SCRIPT_ST= textwrap.dedent("""
    import os, sys
    from PySide6.QtWidgets import QApplication, QWidget, QCheckBox
    from PySide6.QtCore import Qt
    app= QApplication(sys.argv)
    from src.nucleo import bd_single_teacher as db
    from src.nucleo import secuencias_st as S
    from src.nucleo import secuencias as C
    conn= db.obtener_conexion(); c = conn.cursor()
    c.execute("UPDATE configuracion SET nombre_docente='Luis Mejía' WHERE id=1")
    c.execute("INSERT INTO secciones (nombre, nivel_educativo) VALUES ('8vo A','Secundaria')"); sec_id = c.lastrowid
    c.execute("INSERT INTO materias (nombre) VALUES ('Matemática')"); m = c.lastrowid
    c.execute("INSERT INTO materias (nombre) VALUES ('Física')"); m2 = c.lastrowid
    c.execute("INSERT INTO tareas (materia_id, seccion_id, titulo, fecha_entrega) VALUES (?,?, 'Taller', '2026-11-12')", (m, sec_id)); t = c.lastrowid
    c.execute("INSERT INTO tareas (materia_id, titulo) VALUES (?, 'Otra')", (m2,))
    conn.commit(); conn.close()

    #--- lógica
    assert [x["id"] for x in S.tareas_de_materia(m, sec_id)] == [t]
    for malo in ({"titulo": ""}, {"titulo": "x", "fecha_inicio": "2026-12-01", "fecha_fin": "2026-11-01"}):
        try:
            S.guardar_secuencia(malo); raise SystemExit("debió fallar")
        except ValueError:
            pass
    ses= C.sesiones_base(3); ses[0].update(actividad="Explorar", evidencia="Esquema", instrumento="Rúbrica")
    sid= S.guardar_secuencia(
        {"titulo": "Ecuaciones", "materia_id": m, "seccion_id": sec_id,
         "situacion_aprendizaje": "Calculan el costo de una fiesta.",
         "competencias_especificas": " Resuelve ecuaciones \\n\\n Argumenta ", "indicadores_logro": "Plantea la ecuación",
         "competencias_fundamentales": C.unir_competencias_fundamentales(["Resolución de Problemas"]),
         "contenido_conceptos": "Ecuación lineal"}, tareas_ids=[t, t], sesiones=ses)
    sec= S.obtener_secuencia(sid)
    assert sec["materia"] == "Matemática" and sec["grado"] == "8vo A" and sec["profesor"] == "Luis Mejía"
    assert [x["descripcion"] for x in sec["competencias"]] == ["Resuelve ecuaciones", "Argumenta"]
    assert [x["descripcion"] for x in sec["indicadores"]] == ["Plantea la ecuación"]
    assert [x["titulo"] for x in sec["tareas"]] == ["Taller"] and len(sec["sesiones"]) == 3
    avisos= C.validar_secuencia(sec)
    assert any("Sesiones sin actividad" in a for a in avisos) and not any("indicador" in a.lower() for a in avisos), avisos
    assert [r["titulo"] for r in S.listar_secuencias(materia_id=m)] == ["Ecuaciones"]
    assert S.listar_secuencias(materia_id=m2) == [] and S.listar_secuencias(texto="fiesta")[0]["total_sesiones"] == 3
    dup= S.obtener_secuencia(S.duplicar_secuencia(sid))
    assert dup["titulo"].endswith("(copia)") and dup["estado"] == "Borrador" and len(dup["sesiones"]) == 3
    ruta= os.path.join(os.environ["HOME"], "st.docx"); S.exportar_docx(sid, ruta)
    from docx import Document
    d= Document(ruta); texto = "\\n".join(p.text for p in d.paragraphs)
    celdas= {cel.text for tb in d.tables for fila in tb.rows for cel in fila.cells}
    assert "Resuelve ecuaciones" in texto and "Plantea la ecuación" in texto and "Taller" in texto
    assert {"Explorar", "Rúbrica", "Luis Mejía", "8vo A", "Matemática"} <= celdas, celdas
    #borrar materia no borra la secuencia; borrar secuencia no borra tareas
    conn= db.obtener_conexion(); conn.execute("DELETE FROM materias WHERE id=?", (m,)); conn.commit(); conn.close()
    assert S.obtener_secuencia(sid)["materia_id"] is None
    S.eliminar_secuencia(dup["id"]); assert S.obtener_secuencia(dup["id"]) is None

    #--- interfaz
    import src.modulos.secuencias_st as ui
    ui.mostrar_mensaje= lambda *a, **k: None
    mw= QWidget()
    modulo= ui.ModuloSecuenciasST(mw)
    assert modulo.tabla_secuencias.rowCount() == 1
    ed= ui.EditorSecuenciaST(mw, sid, mw)
    assert ed.in_titulo.text() == "Ecuaciones" and ed.in_competencias.toPlainText() == "Resuelve ecuaciones\\nArgumenta"
    assert ed.tabla_sesiones.rowCount() == 3 and ed.combo_curso.currentData() == sec_id
    ed.in_titulo.setText("Ecuaciones II"); ed.tabla_sesiones.cellWidget(0, 8).findChild(QCheckBox).setChecked(True)
    ed._guardar()
    sec= S.obtener_secuencia(sid)
    assert sec["titulo"] == "Ecuaciones II" and C.progreso(sec["sesiones"])[0] == 1
    nuevo= ui.EditorSecuenciaST(mw)
    nuevo.in_titulo.setText("Funciones")
    nuevo.combo_materia.setCurrentIndex(nuevo.combo_materia.findData(m2))
    assert nuevo.lista_tareas.count() == 1           #'Otra'
    nuevo.lista_tareas.item(0).setCheckState(Qt.CheckState.Checked)
    nuevo.in_competencias.setPlainText("Grafica\\nInterpreta")
    nuevo._guardar()
    n= S.obtener_secuencia(nuevo.secuencia_id)
    assert [t["titulo"] for t in n["tareas"]] == ["Otra"] and n["competencias"][1]["descripcion"] == "Interpreta"
    modulo.refrescar(); assert modulo.tabla_secuencias.rowCount() == 2
    modulo.tabla_secuencias.selectRow(0); modulo._duplicar(); assert modulo.tabla_secuencias.rowCount() == 3
    import importlib
    v= importlib.import_module("src.interfaz.ventana_principal")
    mod, cls = v._FABRICAS_MODULOS_SINGLE_TEACHER["Secuencias Didácticas"]
    assert hasattr(importlib.import_module(mod), cls)
    print("OK")
""")


def test_secuencias_single_teacher(tmp_path):

    r= _correr(SCRIPT_ST, tmp_path)
    assert r.returncode == 0 and "OK" in r.stdout, (r.stdout + r.stderr)[-3000:]


SCRIPT_TOUR= textwrap.dedent("""
    import sys
    from PySide6.QtWidgets import QApplication, QWidget, QPushButton
    app= QApplication(sys.argv)
    from src.inicio import tutoriales as T
    from src.interfaz import ventana_principal as V
    nombres= set(V._FABRICAS_MODULOS_SINGLE_TEACHER)
    pasos= [p for p in T.PASOS_SINGLE_TEACHER if p[1] == "Secuencias Didácticas"]
    assert len(pasos) == 1, pasos
    assert all(p[1] is None or p[1] in nombres for p in T.PASOS_SINGLE_TEACHER)
    assert any("Secuencias Didácticas" in c[2] for c in T.CONTENIDO_SINGLE_TEACHER)
    #el texto del botón objetivo debe existir de verdad en la pantalla
    from src.modulos.secuencias_st import ModuloSecuenciasST
    pagina= ModuloSecuenciasST(QWidget()); pagina.show()
    assert T.AsistenteSpotlight._buscar_boton_por_texto(pagina, pasos[0][4]) is not None
    #el orden del tour sigue el del menú: Banco de Tareas -> Secuencias -> Horario
    orden= [p[1] for p in T.PASOS_SINGLE_TEACHER if p[1]]
    assert orden.index("Banco de Tareas") + 1 == orden.index("Secuencias Didácticas") == orden.index("Horario") - 1
    #guardar reporta ÉXITO (es lo único que hace avanzar un paso con acción pendiente)
    import src.modulos.secuencias_st as ui
    from src.nucleo import bd_single_teacher as db
    tipos= []
    ui.mostrar_mensaje= lambda tipo, *a, **k: tipos.append(tipo)
    ed= ui.EditorSecuenciaST(pagina); ed.in_titulo.setText("Tour"); ed._guardar()
    assert tipos == ["exito"], tipos
    print("OK")
""")


def test_tour_single_teacher_incluye_secuencias(tmp_path):

    r= _correr(SCRIPT_TOUR, tmp_path)
    assert r.returncode == 0 and "OK" in r.stdout, (r.stdout + r.stderr)[-3000:]
